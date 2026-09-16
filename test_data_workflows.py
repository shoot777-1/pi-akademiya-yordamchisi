import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
import database as db
import excel_service as excel
import bot


class DataTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.patcher = patch.object(db, "DB_PATH", self.root / "db.sqlite")
        self.patcher.start()
        db.init_db()

    def tearDown(self):
        self.patcher.stop()
        self.tmp.cleanup()

    def workbook(self, invalid=False):
        wb = openpyxl.Workbook()
        pay = wb.active
        pay.title = "oylik to'lovlar"
        pay.cell(5, 3, "Test student")
        pay.cell(5, 4, "Python")
        pay.cell(5, 13, 15)
        sch = wb.create_sheet("Sentabr 2026")
        sch.cell(3, 5, 13)
        sch.cell(5, 3, "Test student")
        sch.cell(5, 4, "Python")
        sch.cell(5, 5, "99:99" if invalid else "14:00")
        path = self.root / "test.xlsx"
        wb.save(path)
        wb.close()
        return path

    def test_dry_run_does_not_change_database(self):
        result = excel.import_pi_academy_excel(self.workbook(), dry_run=True)
        self.assertEqual(result["students_imported"], 1)
        self.assertEqual(db.get_all_students(), [])
        self.assertEqual(db.get_lessons_for_date(date(2026, 9, 13)), [])

    def test_invalid_excel_keeps_previous_data(self):
        db.add_or_update_student("Existing", 2)
        with self.assertRaises(ValueError):
            excel.import_pi_academy_excel(self.workbook(invalid=True))
        self.assertEqual([s["name"] for s in db.get_all_students()], ["Existing"])

    def test_database_failure_rolls_back_student_updates(self):
        db.add_or_update_student("Existing", 2)
        with patch.object(excel, "save_daily_lessons", side_effect=RuntimeError("write failed")):
            with self.assertRaises(RuntimeError):
                excel.import_pi_academy_excel(self.workbook())
        self.assertEqual([s["name"] for s in db.get_all_students()], ["Existing"])

    def test_exact_dates_and_reimport_no_duplicates(self):
        path = self.workbook()
        excel.import_pi_academy_excel(path)
        excel.import_pi_academy_excel(path)
        self.assertEqual(len(db.get_lessons_for_date(date(2026, 9, 13))), 1)
        self.assertEqual(db.get_lessons_for_date(date(2026, 10, 13)), [])
        self.assertEqual(db.get_lessons_for_date(date(2027, 9, 13)), [])
        self.assertEqual(len(db.get_all_students()), 1)

    def test_month_end_due_dates_paid_and_notifications(self):
        student = db.add_or_update_student("Test", 31)
        self.assertEqual(len(db.get_students_due_for_reminder(28, "2027-02-27")), 1)
        db.set_payment(student, "2027-02")
        self.assertEqual(db.get_students_due_for_reminder(28, "2027-02-27"), [])
        db.set_payment(student, "2027-02", False)
        self.assertEqual(len(db.get_students_due_for_reminder(28, "2027-02-27")), 1)
        db.mark_student_notified(student, "2027-02-27")
        self.assertEqual(db.get_students_due_for_reminder(28, "2027-02-27"), [])
        self.assertEqual(len(db.get_students_due_for_reminder(31, "2027-03-30")), 1)
        self.assertEqual(len(db.get_students_due_for_reminder(29, "2028-02-28")), 1)

    def test_cancel_and_move_lesson(self):
        excel.import_pi_academy_excel(self.workbook())
        lesson = db.get_lessons_for_date(date(2026, 9, 13))[0]
        db.change_lesson(lesson["id"])
        self.assertEqual(db.get_lessons_for_date(date(2026, 9, 13)), [])
        db.change_lesson(lesson["id"], "2026-10-01", "15:30")
        moved = db.get_lessons_for_date(date(2026, 10, 1))
        self.assertEqual(moved[0]["lesson_time"], "15:30")
        self.assertEqual(moved[0]["day_number"], 1)

    def test_legacy_paid_values_are_not_confused_with_unpaid(self):
        self.assertTrue(db.is_paid_value("to'landi"))
        self.assertTrue(db.is_paid_value("to’langan"))
        self.assertFalse(db.is_paid_value("to'lanmagan"))
        path = self.workbook()
        wb = openpyxl.load_workbook(path)
        wb["oylik to'lovlar"].cell(5, 9, "to'landi")
        wb.save(path)
        wb.close()
        excel.import_pi_academy_excel(path)
        student = db.get_all_students()[0]
        self.assertIn(student["id"], db.get_paid_students("2026-09"))
        self.assertEqual(db.get_paid_students("2026-10"), set())

    def test_rate_limit_recovers_and_is_per_user(self):
        bot._request_times.clear()
        with patch.object(bot.time, "monotonic", return_value=0):
            self.assertTrue(all(bot.request_allowed(12) for _ in range(6)))
            self.assertFalse(bot.request_allowed(12))
            self.assertTrue(bot.request_allowed(13))
        with patch.object(bot.time, "monotonic", return_value=61):
            self.assertTrue(bot.request_allowed(12))


if __name__ == "__main__":
    unittest.main()
